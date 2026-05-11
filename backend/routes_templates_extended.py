"""
🧩 Templates & Invoice Templates Router (extracted from routes_extended.py)

Domain: Print templates, invoice template designer, XLSX rendering, and public-agent chat.

Extracted on: 2026-02-11
Original location: /app/backend/routes_extended.py lines ~7999-8640

This module keeps the same `/api` URL prefix and behavior — no path changes.
The DB and templates_bucket references are shared with the parent module via
the `set_db(db, bucket)` setter, mimicking the same dependency-injection pattern
used by `routes_extended.set_db`.
"""

from fastapi import APIRouter, HTTPException, Body
from fastapi.responses import HTMLResponse, StreamingResponse
from datetime import datetime, timezone
from typing import Optional, Dict, Any
import io
import os
import uuid

try:
    import openpyxl
except Exception:
    openpyxl = None
try:
    import xlsxwriter
except Exception:
    xlsxwriter = None

from motor.motor_asyncio import AsyncIOMotorGridFSBucket
from supabase_service import SupabaseService

router = APIRouter(prefix="/api", tags=["templates"])

# Module-level shared state (injected by server.py via set_db)
db = None
templates_bucket: Optional[AsyncIOMotorGridFSBucket] = None


def set_db(database, bucket: Optional[AsyncIOMotorGridFSBucket] = None):
    """Inject the shared MongoDB instance and (optionally) a GridFS bucket.

    If bucket is not provided, we initialize one from `database` ourselves
    (same name 'invoice_templates' as the original implementation).
    """
    global db, templates_bucket
    db = database
    if bucket is not None:
        templates_bucket = bucket
        return
    try:
        if db is not None:
            templates_bucket = AsyncIOMotorGridFSBucket(
                db, bucket_name="invoice_templates"
            )
        else:
            templates_bucket = None
    except Exception as e:
        templates_bucket = None
        print(f"templates_extended: GridFS bucket init failed: {e}")


# --------------------- Templates (Compatibility minimal) ---------------------
@router.get("/templates")
async def list_templates():
    # compat: map to invoice templates list
    try:
        docs = await db.invoice_templates.find({}).to_list(length=1000)
        for d in docs:
            d.pop("_id", None)
        return docs
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/templates")
async def create_print_template(payload: Dict[str, Any] = Body(...)):
    """إنشاء نموذج طباعة جديد من المصمم"""
    try:
        doc = {
            "id": str(uuid.uuid4()),
            "type": payload.get("type", "invoice"),
            "name": payload.get("name", "قالب جديد"),
            "content": payload.get("content", ""),
            "isActive": payload.get("isActive", False),
            "createdAt": datetime.now(timezone.utc).isoformat(),
        }
        await db.print_templates.insert_one(doc)
        doc.pop("_id", None)
        return doc
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/templates/{template_id}")
async def delete_template(template_id: str):
    """حذف نموذج طباعة"""
    try:
        result = await db.print_templates.delete_one({"id": template_id})
        if result.deleted_count == 0:
            raise HTTPException(status_code=404, detail="Template not found")
        return {"message": "Template deleted successfully"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/templates/{template_id}/make-default")
async def make_template_default(template_id: str):
    """جعل النموذج افتراضي"""
    try:
        template = await db.print_templates.find_one({"id": template_id})
        if not template:
            raise HTTPException(status_code=404, detail="Template not found")

        template_type = template.get("type", "invoice")

        await db.print_templates.update_many(
            {"type": template_type}, {"$set": {"isActive": False}}
        )

        await db.print_templates.update_one(
            {"id": template_id}, {"$set": {"isActive": True}}
        )

        return {"message": "Template set as default", "id": template_id}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/templates/{template_id}/apply-to-all")
async def apply_template_to_all(template_id: str):
    """تطبيق النموذج على جميع الأنواع"""
    try:
        template = await db.print_templates.find_one({"id": template_id})
        if not template:
            raise HTTPException(status_code=404, detail="Template not found")

        content = template.get("content", "")
        name = template.get("name", "قالب")

        types = ["invoice", "diagnosis", "quote", "receipt", "vehicle_estimate"]
        applied = []

        for t in types:
            await db.print_templates.update_one(
                {"type": t},
                {
                    "$set": {
                        "content": content,
                        "name": f"{name} - {t}",
                        "isActive": True,
                    }
                },
                upsert=True,
            )
            applied.append(t)

        return {"message": "Template applied to all types", "applied": applied}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/print/render", response_class=HTMLResponse)
async def print_render(payload: Dict[str, Any] = Body(...)):
    """Compatibility render endpoint: replaces {{KEY}} placeholders inside HTML."""
    try:
        html = (payload or {}).get("html")
        data = (payload or {}).get("data") or {}
        if not html:
            html = "<html><body><h3>وثيقة</h3><p>{{CUSTOMER_NAME}}</p></body></html>"
        for k, v in data.items():
            html = html.replace(f"{{{{{k}}}}}", str(v))
        return HTMLResponse(content=html)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# --------------------- Print: Resolve Template ---------------------
@router.post("/print/resolve-template")
async def print_resolve_template(payload: Dict[str, Any] = Body(...)):
    try:
        override_type = (
            (payload or {}).get("override_type")
            or (payload or {}).get("type")
            or "invoice"
        )
        tpl = None
        if db:
            tpl = await db.print_templates.find_one(
                {"type": override_type, "isActive": True}
            )
        else:
            supa = SupabaseService()
            templates = await supa.execute_query(
                "print_templates",
                method="select",
                filters={"type": override_type, "is_active": True},
                order=("created_at", "desc"),
                limit=1,
            )
            if templates:
                tpl = templates[0]

        if tpl:
            tpl.pop("_id", None)
            return {"type": override_type, "template": tpl}

        root = os.path.dirname(os.path.abspath(__file__))
        fname = None
        if override_type in ("invoice", "sales_invoice"):
            fname = os.path.join(root, "invoice_template_repair_ar.html")
            if not os.path.exists(fname):
                fname = os.path.join(root, "invoice_template_mechanic.html")
        elif override_type in ("diagnosis", "vehicle_estimate"):
            fname = os.path.join(root, "invoice_template_modern.html")
        elif override_type in ("quote", "receipt"):
            fname = os.path.join(root, "invoice_template_modern.html")
        content = None
        if fname and os.path.exists(fname):
            with open(fname, "r", encoding="utf-8") as f:
                content = f.read()
        else:
            content = "<html><body><h1>قالب طباعة افتراضي</h1><div>{{CUSTOMER_NAME}}</div></body></html>"
        tpl = {
            "id": str(uuid.uuid4()),
            "name": f"Default {override_type}",
            "content": content,
            "isActive": True,
            "type": override_type,
        }
        return {"type": override_type, "template": tpl}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# --------------------- Invoice Templates (Excel-first + Design) ---------------------
@router.get("/invoice-templates")
async def list_invoice_templates():
    provider = os.environ.get("DB_PROVIDER", "mongo").lower()
    if provider in ("memory", "supabase"):
        return []
    docs = (
        await db.invoice_templates.find({"archived": {"$ne": True}})
        .sort("createdAt", -1)
        .to_list(length=1000)
    )
    out = []
    for d in docs:
        d.pop("_id", None)
        out.append(d)
    return out


@router.get("/invoice-templates/{tid}")
async def get_invoice_template(tid: str):
    d = await db.invoice_templates.find_one({"id": tid})
    if not d:
        raise HTTPException(status_code=404, detail="Template not found")
    d.pop("_id", None)
    return d


@router.delete("/invoice-templates/{tid}")
async def delete_invoice_template(tid: str):
    await db.invoice_templates.update_one(
        {"id": tid}, {"$set": {"archived": True, "archivedAt": datetime.utcnow()}}
    )
    return {"status": "archived"}


@router.put("/invoice-templates/{tid}")
async def update_invoice_template(tid: str, payload: Dict[str, Any] = Body(...)):
    await db.invoice_templates.update_one(
        {"id": tid}, {"$set": {**payload, "updatedAt": datetime.utcnow()}}
    )
    doc = await db.invoice_templates.find_one({"id": tid})
    if not doc:
        raise HTTPException(status_code=404, detail="Template not found")
    doc.pop("_id", None)
    return doc


@router.post("/invoice-templates/create-blank")
async def create_blank_template(payload: Dict[str, Any] = Body(None)):
    try:
        name = (payload or {}).get("name") or "قالب فارغ"
        rows = int((payload or {}).get("rows") or 12)
        cols = int((payload or {}).get("cols") or 8)
        preview = [["" for _ in range(cols)] for _ in range(rows)]
        if rows >= 2:
            preview[0][0] = "{{WORKSHOP_NAME}}"
            preview[0][3] = "{{CUSTOMER_NAME}}"
            preview[1][0] = "{{ITEMS}}"
        doc = {
            "id": str(uuid.uuid4()),
            "name": name,
            "format": "xlsx",
            "fileId": None,
            "fields": ["{{WORKSHOP_NAME}}", "{{CUSTOMER_NAME}}", "{{ITEMS}}"],
            "preview": preview,
            "mapping": {},
            "itemsConfig": {
                "anchor": "{{ITEMS}}",
                "columns": {"description": "", "qty": "", "price": "", "total": ""},
            },
            "elements": [],
            "schema": [],
            "page": {"size": "A4", "orientation": "portrait"},
            "isDefault": False,
            "createdAt": datetime.utcnow(),
        }
        await db.invoice_templates.insert_one(doc)
        doc.pop("_id", None)
        return doc
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/invoice-templates/{tid}/design")
async def save_design(tid: str, payload: Dict[str, Any] = Body(...)):
    try:
        elements = (payload or {}).get("elements") or []
        schema = (payload or {}).get("schema") or []
        page = (payload or {}).get("page") or {"size": "A4", "orientation": "portrait"}
        await db.invoice_templates.update_one(
            {"id": tid},
            {
                "$set": {
                    "elements": elements,
                    "schema": schema,
                    "page": page,
                    "updatedAt": datetime.utcnow(),
                }
            },
        )
        d = await db.invoice_templates.find_one({"id": tid})
        if not d:
            raise HTTPException(status_code=404, detail="Template not found")
        d.pop("_id", None)
        return d
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/invoice-templates/{tid}/save-json")
async def save_template_from_json(tid: str, payload: Dict[str, Any] = Body(...)):
    alt_mapping = (payload or {}).get("mapping")
    alt_items = (payload or {}).get("items")
    if alt_mapping is not None or alt_items is not None:
        try:
            await db.invoice_templates.update_one(
                {"id": tid},
                {
                    "$set": {
                        "mapping": alt_mapping or {},
                        "itemsConfig": alt_items or {},
                        "updatedAt": datetime.utcnow(),
                    }
                },
            )
        except Exception:
            pass
    try:
        if not xlsxwriter:
            raise HTTPException(status_code=500, detail="xlsxwriter not installed")
        grid = payload.get("grid")
        if not isinstance(grid, list):
            raise HTTPException(status_code=400, detail="grid required")
        out = io.BytesIO()
        book = xlsxwriter.Workbook(out, {"in_memory": True})
        sheet = book.add_worksheet("Template")
        for r, row in enumerate(grid):
            if not isinstance(row, list):
                continue
            for c, val in enumerate(row):
                sheet.write(r, c, "" if val is None else str(val))
        book.close()
        xlsx_bytes = out.getvalue()
        file_id = None
        if templates_bucket:
            try:
                file_oid = await templates_bucket.upload_from_stream(
                    f"template_{tid}.xlsx",
                    io.BytesIO(xlsx_bytes),
                    metadata={
                        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    },
                )
                file_id = str(file_oid)
            except Exception as e:
                print(f"gridfs upload failed: {e}")
        await db.invoice_templates.update_one(
            {"id": tid},
            {
                "$set": {
                    "format": "xlsx",
                    "fileId": file_id,
                    "updatedAt": datetime.utcnow(),
                }
            },
        )
        return {"status": "ok", "fileId": file_id}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/invoice-templates/{tid}/make-default")
async def make_default_template(tid: str):
    try:
        await db.invoice_templates.update_many({}, {"$set": {"isDefault": False}})
        await db.invoice_templates.update_one(
            {"id": tid}, {"$set": {"isDefault": True, "updatedAt": datetime.utcnow()}}
        )
        doc = await db.invoice_templates.find_one({"id": tid})
        if not doc:
            raise HTTPException(status_code=404, detail="Template not found")
        doc.pop("_id", None)
        return doc
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/invoice-templates/{tid}/update-mapping")
async def update_template_mapping(tid: str, payload: Dict[str, Any] = Body(...)):
    try:
        mapping = (payload or {}).get("mapping") or {}
        items = (payload or {}).get("items") or {}
        await db.invoice_templates.update_one(
            {"id": tid},
            {
                "$set": {
                    "mapping": mapping,
                    "itemsConfig": items,
                    "updatedAt": datetime.utcnow(),
                }
            },
        )
        doc = await db.invoice_templates.find_one({"id": tid})
        if not doc:
            raise HTTPException(status_code=404, detail="Template not found")
        doc.pop("_id", None)
        return doc
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/invoice-templates/{tid}/save-named")
async def save_named_copy(tid: str, payload: Dict[str, Any] = Body(...)):
    try:
        if not xlsxwriter:
            raise HTTPException(status_code=500, detail="xlsxwriter not installed")
        name = (payload or {}).get("name") or "فاتوره"
        grid = (payload or {}).get("grid") or []
        mapping = (payload or {}).get("mapping") or {}
        items_cfg = (payload or {}).get("itemsConfig") or {}
        sample_data = (payload or {}).get("data") or {}
        elements = (payload or {}).get("elements") or []
        schema = (payload or {}).get("schema") or []
        page = (payload or {}).get("page") or {"size": "A4", "orientation": "portrait"}

        out = io.BytesIO()
        book = xlsxwriter.Workbook(out, {"in_memory": True})
        sheet = book.add_worksheet("Template")
        for r, row in enumerate(grid):
            if not isinstance(row, list):
                continue
            for c, val in enumerate(row):
                sheet.write(r, c, "" if val is None else str(val))
        book.close()
        xlsx_bytes = out.getvalue()
        file_id = None
        if templates_bucket:
            try:
                file_oid = await templates_bucket.upload_from_stream(
                    f"{name}_{uuid.uuid4().hex}.xlsx",
                    io.BytesIO(xlsx_bytes),
                    metadata={
                        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    },
                )
                file_id = str(file_oid)
            except Exception as e:
                print(f"gridfs upload failed: {e}")
        new_id = str(uuid.uuid4())
        doc = {
            "id": new_id,
            "name": name,
            "format": "xlsx",
            "fileId": file_id,
            "fields": [],
            "preview": grid,
            "mapping": mapping,
            "itemsConfig": items_cfg,
            "elements": elements,
            "schema": schema,
            "page": page,
            "sampleData": sample_data,
            "isDefault": False,
            "createdAt": datetime.utcnow(),
        }
        await db.invoice_templates.insert_one(doc)
        doc.pop("_id", None)
        return doc
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/invoice-templates/{tid}/auto-save")
async def auto_save_template(tid: str, payload: Dict[str, Any] = Body(...)):
    """Auto-save from Studio: persist design/grid/mapping/items + ensure snapshot copy."""
    try:
        grid = (payload or {}).get("grid")
        mapping = (payload or {}).get("mapping")
        items_cfg = (payload or {}).get("itemsConfig")
        elements = (payload or {}).get("elements")
        schema = (payload or {}).get("schema")
        page = (payload or {}).get("page")
        updates = {"updatedAt": datetime.utcnow()}
        if grid is not None:
            updates["preview"] = grid
        if mapping is not None:
            updates["mapping"] = mapping
        if items_cfg is not None:
            updates["itemsConfig"] = items_cfg
        if elements is not None:
            updates["elements"] = elements
        if schema is not None:
            updates["schema"] = schema
        if page is not None:
            updates["page"] = page
        if updates:
            await db.invoice_templates.update_one({"id": tid}, {"$set": updates})
        snap = await db.invoice_templates.find_one(
            {"name": {"$regex": "^فاتوره"}, "sourceId": tid, "archived": {"$ne": True}}
        )
        if not snap:
            new_id = str(uuid.uuid4())
            base = await db.invoice_templates.find_one({"id": tid})
            if base:
                base.pop("_id", None)
                base.update(
                    {
                        "id": new_id,
                        "name": "فاتوره",
                        "sourceId": tid,
                        "isDefault": False,
                        "createdAt": datetime.utcnow(),
                        "updatedAt": datetime.utcnow(),
                    }
                )
                await db.invoice_templates.insert_one(base)
        doc = await db.invoice_templates.find_one({"id": tid})
        doc.pop("_id", None)
        return {"status": "ok", "template": doc}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/print/invoice-xlsx")
async def print_invoice_xlsx(payload: Dict[str, Any] = Body(...)):
    try:
        if not openpyxl:
            raise HTTPException(status_code=500, detail="openpyxl not installed")
        t_id = payload.get("templateId")
        data = payload.get("data") or {}

        template = await db.invoice_templates.find_one({"id": t_id})
        if not template:
            raise HTTPException(status_code=404, detail="Template not found")

        out = io.BytesIO()
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Invoice"

        sheet["A1"] = data.get("WORKSHOP_NAME", "Workshop")
        sheet["A2"] = f"Invoice: {data.get('INVOICE_NO', '')}"

        row = 5
        sheet.cell(row, 1, "Item")
        sheet.cell(row, 2, "Qty")
        sheet.cell(row, 3, "Price")
        sheet.cell(row, 4, "Total")

        items = data.get("ITEMS", [])
        for item in items:
            row += 1
            sheet.cell(row, 1, item.get("description", ""))
            sheet.cell(row, 2, item.get("qty", 0))
            sheet.cell(row, 3, item.get("price", 0))
            sheet.cell(row, 4, item.get("total", 0))

        book.save(out)
        out.seek(0)

        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=invoice_{data.get('INVOICE_NO')}.xlsx"
            },
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ============ Public Agent Proxy (Internal LLM) ============
@router.post("/public-agent/chat")
async def public_agent_chat(payload: Dict[str, Any] = Body(...)):
    """Public-facing chat using EMERGENT_LLM_KEY (Claude Sonnet)."""
    try:
        message = payload.get("message")
        session_id = payload.get("sessionId") or str(uuid.uuid4())

        llm_key = os.getenv("EMERGENT_LLM_KEY")
        if not llm_key:
            return {
                "response": "عذراً، خدمة المحادثة غير متوفرة حالياً (API Key missing).",
                "session_id": session_id,
            }

        # Local import to avoid loading the SDK at module init (heavy)
        from emergentintegrations.llm.chat import LlmChat, UserMessage

        system_message = """أنت مساعد ذكي لورشة سيارات. تتحدث العربية بطلاقة.
        مهمتك مساعدة العملاء في الإجابة على استفساراتهم حول صيانة السيارات، المواعيد، والخدمات.
        كن مهذباً ومحترفاً."""

        chat = LlmChat(
            api_key=llm_key, session_id=session_id, system_message=system_message
        ).with_model("anthropic", "claude-sonnet-4.5-20250929")

        response_text = await chat.send_message(UserMessage(text=message))

        return {"response": response_text, "session_id": session_id}
    except Exception as e:
        print(f"Public Agent Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
